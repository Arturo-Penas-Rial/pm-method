#!/usr/bin/env python3
"""Exporta artefactos Markdown a .docx y registros CSV a .xlsx.
Export Markdown artifacts to .docx and CSV registers to .xlsx.

Requiere: pip install python-docx openpyxl   (si no están, el script lo indica y sale).
Uso / Usage:
  python3 exportar.py docx --entrada 01-inicio/acta-constitucion.md --salida exportados/acta-constitucion.docx
  python3 exportar.py xlsx --proyecto ./mi-proyecto            # todos los registros en un libro, una hoja por registro
  python3 exportar.py xlsx --entrada 03-control/registros/riesgos.csv --salida exportados/riesgos.xlsx
"""
import argparse, csv, os, re, sys, glob

def md_a_docx(entrada, salida):
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        sys.exit("Falta python-docx: pip install python-docx")
    doc = Document(); txt = open(entrada, encoding="utf-8").read()
    estado_doc = ""
    if txt.startswith("---"):
        parts = txt.split("---", 2)
        if len(parts) == 3:
            m = re.search(r"(?m)^estado:\s*\"?([\w-]+)", parts[1]); estado_doc = m.group(1) if m else ""
            txt = parts[2]
    if estado_doc != "aprobado":
        p = doc.add_paragraph(); r = p.add_run(f"BORRADOR / DRAFT — estado: {estado_doc or 'sin estado'} · generado con asistencia de IA, requiere revisión humana")
        r.bold = True
    import shutil, subprocess, tempfile
    mmdc = shutil.which("mmdc")
    bloque, en_bloque, lenguaje = [], False, ""
    def flush_bloque():
        codigo = "\n".join(bloque)
        if lenguaje == "mermaid" and mmdc:
            with tempfile.TemporaryDirectory() as td:
                src, png = os.path.join(td, "d.mmd"), os.path.join(td, "d.png")
                open(src, "w", encoding="utf-8").write(codigo)
                if subprocess.run([mmdc, "-i", src, "-o", png, "-w", "1600"], capture_output=True).returncode == 0 and os.path.exists(png):
                    from docx.shared import Inches
                    doc.add_picture(png, width=Inches(6.3)); return
        if lenguaje == "mermaid":
            doc.add_paragraph("[Diagrama Mermaid — renderizar con pm-visuales o mermaid-cli (mmdc); fuente:]").runs[0].italic = True
        par = doc.add_paragraph(); run = par.add_run(codigo); run.font.name = "Courier New"; run.font.size = Pt(8)
    tabla = []
    def flush_tabla():
        nonlocal tabla
        if not tabla: return
        filas = [r for r in tabla if not re.match(r"^\|?\s*:?-{2,}", r)]
        celdas = [[c.strip() for c in r.strip().strip("|").split("|")] for r in filas]
        if celdas:
            t = doc.add_table(rows=len(celdas), cols=max(len(c) for c in celdas)); t.style = "Table Grid"
            for i, fila in enumerate(celdas):
                for j, c in enumerate(fila): t.cell(i, j).text = c
        tabla = []
    for line in txt.splitlines():
        if line.strip().startswith("```"):
            if en_bloque:
                flush_bloque(); bloque, en_bloque = [], False
            else:
                flush_tabla(); en_bloque, lenguaje = True, line.strip()[3:].strip().lower()
            continue
        if en_bloque: bloque.append(line); continue
        if line.strip().startswith("|"): tabla.append(line); continue
        flush_tabla()
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m: doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4)); continue
        if re.match(r"^\s*[-*]\s+\[[ xX]\]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet"); continue
        if re.match(r"^\s*[-*]\s+", line): doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet"); continue
        if re.match(r"^\s*\d+\.\s+", line): doc.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", line), style="List Number"); continue
        if line.strip() == "": continue
        p = doc.add_paragraph()
        for seg in re.split(r"(\*\*[^*]+\*\*)", line):
            if seg.startswith("**") and seg.endswith("**"): p.add_run(seg[2:-2]).bold = True
            else: p.add_run(seg)
    flush_tabla()
    os.makedirs(os.path.dirname(salida) or ".", exist_ok=True); doc.save(salida); print(f"Escrito {salida}")

def csv_a_xlsx(entradas, salida):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")
    wb = Workbook(); wb.remove(wb.active)
    for e in entradas:
        ws = wb.create_sheet(os.path.splitext(os.path.basename(e))[0][:31])
        for i, row in enumerate(csv.reader(open(e, encoding="utf-8"))):
            ws.append(row)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="002855")
        for col in range(1, ws.max_column + 1): ws.column_dimensions[get_column_letter(col)].width = 18
        ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(salida) or ".", exist_ok=True); wb.save(salida); print(f"Escrito {salida}")

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("formato", choices=["docx", "xlsx"])
    ap.add_argument("--entrada"); ap.add_argument("--salida"); ap.add_argument("--proyecto")
    a = ap.parse_args()
    if a.formato == "docx":
        if not (a.entrada and a.salida): sys.exit("docx requiere --entrada y --salida")
        md_a_docx(a.entrada, a.salida)
    else:
        if a.proyecto:
            ents = sorted(glob.glob(os.path.join(a.proyecto, "03-control", "registros", "*.csv")))
            import datetime
            csv_a_xlsx(ents, a.salida or os.path.join(a.proyecto, "exportados", f"{datetime.date.today().isoformat()}-registros.xlsx"))
        elif a.entrada and a.salida: csv_a_xlsx([a.entrada], a.salida)
        else: sys.exit("xlsx requiere --proyecto o --entrada y --salida")

if __name__ == "__main__":
    main()
